"""Offset-based reports: preserve text, highlight only positive decisions, expose unknowns."""
from copy import deepcopy
from html import escape
from pathlib import Path
from urllib.parse import urlparse
from src.document_processing.file_reader import iter_docx_paragraphs,read_document


def _ranges(result):
    text=result['original_text'];ranges=[]
    for row in sorted(result['results'],key=lambda r:r['start']):
        start,end=row['start'],row['end']
        if not 0<=start<end<=len(text):raise ValueError('Invalid sentence offsets')
        if ' '.join(text[start:end].split())!=row['sentence']:raise ValueError('Sentence/offset mismatch')
        if row['status']=='suspicious':
            if ranges and start<ranges[-1][1]:raise ValueError('Overlapping highlight spans')
            ranges.append((start,end,row['sentence_id']))
    return ranges


def _details(row):
    pieces=[f"Câu {row['sentence_id']+1}: {row['label']}"]
    if row.get('probability') is not None:pieces.append(f"Điểm dự đoán (chưa hiệu chỉnh): {row['probability']:.4f}")
    if row.get('source_url'):pieces.append('Nguồn đối chiếu: '+row['source_url'])
    if row.get('matched_source_sentence'):pieces.append('Đoạn nguồn: '+row['matched_source_sentence'])
    if row.get('explanation'):pieces.append('Giải thích: '+row['explanation'])
    if row.get('explanation_status')=='failed':pieces.append('Tầng 5 không tạo được giải thích; kết quả chưa đầy đủ.')
    if row.get('complete') is False:pieces.append('Chưa hoàn tất tất cả nguồn hoặc tầng xử lý.')
    if row.get('errors'):pieces.append('Lỗi: '+str(row['errors']))
    return '\n'.join(pieces)


def export_highlighted_html(process_result,output_path,original_sentences=None):
    text=process_result['original_text'];ranges=_ranges(process_result)
    fragments=[];previous=0
    for start,end,index in ranges:
        fragments.extend([escape(text[previous:start]),f'<mark id="sentence-{index}">{escape(text[start:end])}</mark>'])
        previous=end
    fragments.append(escape(text[previous:]))
    rows=[]
    for row in process_result['results']:
        link=''
        url=row.get('source_url')
        if url and urlparse(url).scheme in ('https','http'):
            link=f'<a href="{escape(url,quote=True)}" rel="noopener noreferrer">Mở nguồn</a>'
        rows.append('<tr><td>'+str(row['sentence_id']+1)+'</td><td>'+escape(row['sentence'])+'</td><td class="details">'+escape(_details(row))+' '+link+'</td></tr>')
    warnings=''.join('<li>'+escape(w)+'</li>' for w in process_result.get('warnings',[]))
    html='''<!doctype html><html lang="vi"><meta charset="utf-8"><title>Báo cáo đối chiếu</title>
<style>body{font:16px/1.6 sans-serif;max-width:1100px;margin:32px auto;padding:16px}.document,.details{white-space:pre-wrap}mark{background:#ffe49c}table{border-collapse:collapse;width:100%;margin-top:24px}td,th{border:1px solid #ccc;padding:10px;vertical-align:top;overflow-wrap:anywhere}td:first-child{width:40px}</style>'''
    html+=f"<h1>Báo cáo: {escape(Path(process_result['file_path']).name)}</h1><p>Số câu: {process_result['total_sentences']} · Nghi vấn: {process_result['suspicious_sentences']}</p>"
    html+='<p>Highlight biểu thị nghi vấn cần đối chiếu, không phải kết luận đạo văn. Các trạng thái chưa kiểm tra được liệt kê bên dưới.</p><ul>'+warnings+'</ul>'
    html+='<div class="document">'+''.join(fragments)+'</div><h2>Chi tiết từng câu</h2><table><tr><th>#</th><th>Câu</th><th>Kết quả và nguồn</th></tr>'+''.join(rows)+'</table></html>'
    output=Path(output_path);output.parent.mkdir(parents=True,exist_ok=True);output.write_text(html,encoding='utf-8')
    return output


def _highlight_paragraph(paragraph,spans):
    """Split XML runs while keeping formatting, hyperlinks and non-text children."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    def child_text(child):
        if child.tag==qn('w:t'):return child.text or ''
        if child.tag==qn('w:tab'):return '\t'
        if child.tag==qn('w:cr'):return '\n'
        if child.tag==qn('w:br') and child.get(qn('w:type'),'textWrapping')=='textWrapping':return '\n'
        return ''
    runs=[]
    for run in paragraph._p.iter(qn('w:r')):
        nearest=next((p for p in run.iterancestors() if p.tag==qn('w:p')),None)
        if nearest is paragraph._p:runs.append(run)
    if ''.join(child_text(c) for run in runs for c in run)!=paragraph.text:
        raise ValueError('Unsupported Word paragraph structure; refusing misplaced highlights')
    offset=0
    for run in runs:
        children=[c for c in run if c.tag!=qn('w:rPr')]
        length=sum(len(child_text(c)) for c in children)
        if not any(a<offset+length and b>offset for a,b in spans):offset+=length;continue
        parent=run.getparent();position=parent.index(run);properties=run.find(qn('w:rPr'))
        for child in children:
            content=child_text(child)
            boundaries={0,len(content)}
            for a,b in spans:
                boundaries.update([max(0,min(len(content),a-offset)),max(0,min(len(content),b-offset))])
            points=sorted(boundaries)
            chunks=list(zip(points,points[1:])) or [(0,0)]
            for a,b in chunks:
                new=OxmlElement('w:r')
                for key,value in run.attrib.items():new.set(key,value)
                props=deepcopy(properties) if properties is not None else OxmlElement('w:rPr')
                highlighted=any(x<offset+b and y>offset+a for x,y in spans)
                if highlighted:
                    for old in list(props.findall(qn('w:highlight'))):props.remove(old)
                    highlight=OxmlElement('w:highlight');highlight.set(qn('w:val'),'yellow');props.append(highlight)
                if len(props):new.append(props)
                element=deepcopy(child)
                if child.tag==qn('w:t'):
                    element.text=content[a:b];element.set(qn('xml:space'),'preserve')
                new.append(element);parent.insert(position,new);position+=1
            offset+=len(content)
        parent.remove(run)


def export_highlighted_docx(process_result,output_path,original_sentences=None):
    import docx
    from docx.enum.text import WD_COLOR_INDEX
    source=Path(process_result['file_path']);output=Path(output_path)
    if source.resolve()==output.resolve():raise ValueError('Choose an output path different from the source')
    text=process_result['original_text'];ranges=_ranges(process_result)
    if process_result.get('format')=='docx':
        if read_document(source)['text']!=text:raise ValueError('Source document changed since analysis')
        document=docx.Document(source);offset=0
        for paragraph in list(iter_docx_paragraphs(document)):
            end=offset+len(paragraph.text)
            spans=[(max(a,offset)-offset,min(b,end)-offset) for a,b,_ in ranges if a<end and b>offset]
            if spans:_highlight_paragraph(paragraph,spans)
            offset=end+1
    else:
        document=docx.Document();paragraph=document.add_paragraph();previous=0
        for start,end,_ in ranges:
            paragraph.add_run(text[previous:start]);paragraph.add_run(text[start:end]).font.highlight_color=WD_COLOR_INDEX.YELLOW
            previous=end
        paragraph.add_run(text[previous:])
    document.add_page_break();document.add_heading('Kết quả đối chiếu',level=1)
    document.add_paragraph(f"Số câu: {process_result['total_sentences']} | Nghi vấn: {process_result['suspicious_sentences']}")
    document.add_paragraph('Highlight biểu thị nghi vấn cần đối chiếu; không phải kết luận đạo văn.')
    for warning in process_result.get('warnings',[]):document.add_paragraph(warning)
    for row in process_result['results']:document.add_paragraph(_details(row))
    output.parent.mkdir(parents=True,exist_ok=True);document.save(output)
    return output
