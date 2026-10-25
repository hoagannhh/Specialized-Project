import ast,json,sys,types
from pathlib import Path
sys.stdout.reconfigure(encoding='utf-8')
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from src.document_processing.sentence_splitter import split_sentences
from src.document_processing.highlight_exporter import export_highlighted_html
root=Path('tmp/source-review');root.mkdir(parents=True,exist_ok=True)
p=Path('src/pipeline/single_file_pipeline.py')
nodes=[n for n in ast.parse(p.read_text(encoding='utf-8-sig')).body if isinstance(n,(ast.FunctionDef,ast.ClassDef))]
text='A sufficiently long sentence for checking.'
ns={'Path':Path,'TANG0_CHECKPOINT_DIR':'unused','WEB_SEARCH_MAX_RESULTS':3,'split_sentences':split_sentences,'read_file':lambda p:text,'logger':types.SimpleNamespace(info=lambda *a:None)}
exec(compile(ast.Module(body=nodes,type_ignores=[]),str(p),'exec'),ns)
obj=ns['SingleFilePipeline'].__new__(ns['SingleFilePipeline'])
obj.tang0_checkpoint_dir='unused';obj.max_search_results=3
obj._tang0_inference=types.SimpleNamespace(predict_sentence=lambda *a,**k:{'decision':'suspicious','probability':0.9})
obj._web_search=types.SimpleNamespace(search_web=lambda *a,**k:[{'url':'https://example.org'}])
obj._extract_page_text=types.SimpleNamespace(fetch_and_extract_text=lambda *a:{'success':True,'text':'Another sufficiently long source sentence.'})
obj.cascade_pipeline=types.SimpleNamespace(predict=lambda *a:{'label':'không đạo văn','probability':0.01})
r=obj.process_file('dummy.txt')
html=export_highlighted_html(r,str(root/'negative.html'),[text]).read_text(encoding='utf-8')
print('NEGATIVE_RESULT',json.dumps({'label':r['results'][0]['label'],'suspicious_count':r['suspicious_sentences'],'highlighted':"<span class='low'>" in html},ensure_ascii=False))
obj._web_search.search_web=lambda *a,**k:[]
r=obj.process_file('dummy.txt')
html=export_highlighted_html(r,str(root/'unknown.html'),[text]).read_text(encoding='utf-8')
print('UNKNOWN_STATUS_VISIBLE',r['results'][0]['label'] in html)
print('SHORT_TEXT',split_sentences('Title. Short. This is a sufficiently long paragraph.'))
find=ns['find_best_matching_sentence']
print('ZERO_OVERLAP_CANDIDATE',find('Quantum entanglement experiments.','Bananas grow beside tropical orchards.'))
from src.models.tang5_llm_judge import prompts
p=Path('src/models/tang5_llm_judge/call_api.py')
f=next(n for n in ast.parse(p.read_text(encoding='utf-8-sig')).body if isinstance(n,ast.FunctionDef) and n.name=='_extract_json_from_text')
ns2={'json':json,'re':__import__('re')}
exec(compile(ast.Module(body=[f],type_ignores=[]),str(p),'exec'),ns2)
print('LLM_JSON_WRONG_SCHEMA',type(ns2['_extract_json_from_text']('[]')).__name__)
try:
 import docx
 from src.document_processing.file_reader import read_file
 d=docx.Document();d.add_paragraph('Body paragraph.');d.add_table(rows=1,cols=1).cell(0,0).text='Important table text.'
 path=root/'table.docx';d.save(path)
 print('TABLE_TEXT_PRESERVED','Important table text.' in read_file(str(path)))
except ImportError as e:
 print('DOCX_TEST_UNAVAILABLE',str(e))
