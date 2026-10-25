from pathlib import Path
import sys
sys.stdout.reconfigure(encoding='utf-8')
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
import textwrap

ROOT=Path(__file__).resolve().parent
OUT=ROOT.parents[1]/'output'/'documents'
OUT.mkdir(parents=True,exist_ok=True)
fontpath='C:/Windows/Fonts/arial.ttf'
font=ImageFont.truetype(fontpath,30)
small=ImageFont.truetype(fontpath,26)
def diagram(name,lines):
    w=1500; h=len(lines)*130+25
    im=Image.new('RGB',(w,h),'white'); d=ImageDraw.Draw(im)
    for i,label in enumerate(lines):
        y=15+i*130
        d.rounded_rectangle((80,y,w-80,y+90),radius=12,fill='#f3f5f7',outline='#5b6770',width=2)
        chunks=label.split('|')
        for j,t in enumerate(chunks):
            box=d.textbbox((0,0),t,font=font)
            d.text(((w-(box[2]-box[0]))/2,y+(20 if len(chunks)==1 else 8)+j*36),t,font=font,fill='black')
        if i<len(lines)-1:
            d.line((750,y+93,750,y+122),fill='#485560',width=3)
            d.polygon([(750,y+127),(740,y+115),(760,y+115)],fill='#485560')
    path=ROOT/(name+'.png'); im.save(path); return path

figs={
'overall':diagram('overall',[
'Tệp đầu vào → Trích xuất và lưu vị trí gốc',
'Tách câu và giữ ngữ cảnh → Bước 0 chấm dấu hiệu AI',
'Đầy đủ: mọi đơn vị đủ thông tin | Tiết kiệm: ưu tiên theo điểm và ngân sách',
'Tìm nguồn → Đọc trang → Chia đoạn → Lấy hợp ứng viên',
'Đối chiếu bằng CaTE-X → Tổng hợp nguồn và trích dẫn',
'Xuất highlight và bảng bằng chứng|Phân biệt chưa kiểm tra, không có nguồn và nguồn không đọc được']),
'cascade':diagram('cascade',[
'Cặp nguồn A và đoạn cần kiểm tra B',
'Tầng 1 TF-IDF → Tầng 2 SBERT pretrained',
'Chỉ lọc nếu ở chế độ tiết kiệm và cả hai điểm thấp|Các trường hợp còn lại chuyển tầng 3',
'Tầng 3 Bi-Encoder Triplet|Điểm thấp: chưa đủ bằng chứng; điểm cao: nghi vấn; vùng giữa: tầng 4',
'Tầng 4 Cross-Encoder cho cặp khó|Quyết định theo ngưỡng validation và giữ vùng cần xem xét',
'Tầng 5 giải thích các cặp cần báo cáo từ tầng 3 hoặc 4|Không đổi nhãn, không tự tạo nguồn']),
'split':diagram('split',[
'Nhóm nguồn và toàn bộ paraphrase liên quan',
'Rà soát trùng và gần trùng → Gán group_id',
'Chia nhóm train / validation / test một lần',
'Trong mỗi split: tạo đơn câu, cặp và triplet|Negative chỉ được lấy từ chính split đó',
'Dùng chung bản kê chia tập cho tất cả mô hình và baseline'])}

doc=Document(); sec=doc.sections[0]
sec.page_height=Cm(29.7);sec.page_width=Cm(21)
sec.top_margin=Cm(2.3);sec.bottom_margin=Cm(2.2);sec.left_margin=Cm(3);sec.right_margin=Cm(2)
normal=doc.styles['Normal']; normal.font.name='Times New Roman';normal.font.size=Pt(13)
normal.paragraph_format.line_spacing=1.3
normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.first_line_indent=Cm(.75)
normal.paragraph_format.widow_control=True
for nm,size in [('Title',17),('Heading 1',15),('Heading 2',13),('Caption',11)]:
    s=doc.styles[nm];s.font.name='Times New Roman';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.first_line_indent=Cm(0)
    s.paragraph_format.space_before=Pt(12 if nm!='Title' else 0)
    s.paragraph_format.space_after=Pt(7)
    if nm!='Caption':s.font.bold=True
    s.paragraph_format.keep_with_next=True
    s.paragraph_format.line_spacing=1.15
doc.styles['Caption'].font.italic=True
for s in doc.styles:
    for node in list(s.element.iter(qn('w:pBdr'))):node.getparent().remove(node)
    if s.name in ['Normal','Title','Heading 1','Heading 2','Caption']:
        rf=s.element.get_or_add_rPr().find(qn('w:rFonts'))
        if rf is not None:
            for key in list(rf.attrib):
                if key.endswith('Theme'):del rf.attrib[key]
            for key in ['ascii','hAnsi','eastAsia','cs']:rf.set(qn('w:'+key),'Times New Roman')
# Page number is local to the standalone chapter.
p=sec.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.first_line_indent=Cm(0)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');p._p.append(fld)
doc.core_properties.title='Chương 2 Phương pháp đề xuất CaTE-X'
doc.core_properties.subject='Phát hiện AI-paraphrase với Triplet Loss và bằng chứng nguồn'
doc.core_properties.author='Nguyễn Hồng An; Nguyễn Hoàng Anh'

def para(txt):
    p=doc.add_paragraph(txt);p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    if txt.startswith('['):
        p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1.1
        for r in p.runs:r.font.size=Pt(11)
    return p
def table(caption,rows):
    p=doc.add_paragraph(caption,'Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    t=doc.add_table(rows=1,cols=len(rows[0]));t.autofit=False
    widths=[3.1,4.7,8.2] if len(rows[0])==3 else [5,11]
    for c,w in zip(t.columns,widths):c.width=Cm(w)
    for idx,row in enumerate(rows):
        cells=t.rows[0].cells if idx==0 else t.add_row().cells
        for c,txt,w in zip(cells,row,widths):
            c.width=Cm(w);c.text=txt
            tcPr=c._tc.get_or_add_tcPr()
            borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            tcPr.append(borders)
            mar=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),'90');el.set(qn('w:type'),'dxa');mar.append(el)
            tcPr.append(mar)
            if idx==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E8ECF0');tcPr.append(sh)
            for p in c.paragraphs:
                p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1.1;p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
                p.paragraph_format.keep_with_next=(idx==0 or idx==len(rows)-2)
                for r in p.runs:r.font.size=Pt(11);r.bold=idx==0
        pr=t.rows[idx]._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        if idx==0:pr.append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after=Pt(0)

lines=(ROOT/'chapter.md').read_text(encoding='utf-8').splitlines();i=0
while i<len(lines):
    line=lines[i].strip();i+=1
    if not line:continue
    if line.startswith('# '):
        p=doc.add_paragraph(line[2:],'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    elif line.startswith('## '):doc.add_paragraph(line[3:],'Heading 1')
    elif line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2')
    elif line.startswith('@figure '):
        name,cap=line[8:].split('|',1)
        p=doc.add_paragraph();p.paragraph_format.first_line_indent=Cm(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
        r=p.add_run();r.add_picture(str(figs[name]),width=Cm(15.7))
        p=doc.add_paragraph(cap,'Caption');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=False
    elif line.startswith('@table '):
        rows=[]
        while i<len(lines) and lines[i]!='@end':rows.append(lines[i].split('|'));i+=1
        i+=1;table(line[7:],rows)
    elif line.startswith('@eq '):
        p=doc.add_paragraph();p.paragraph_format.first_line_indent=Cm(0);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        math=OxmlElement('m:oMath');mr=OxmlElement('m:r');mt=OxmlElement('m:t');mt.text=line[4:];mr.append(mt);math.append(mr);p._p.append(math)
    else:para(line)
doc.save(OUT/'Chuong_2_Phuong_phap_CaTE-X.docx')
print(OUT/'Chuong_2_Phuong_phap_CaTE-X.docx')
