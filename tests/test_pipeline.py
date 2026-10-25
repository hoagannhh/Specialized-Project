import unittest
from pathlib import Path
from uuid import uuid4
from unittest.mock import patch
import numpy as np
import pandas as pd
from src.pipeline.cascade_pipeline import CaTEXPipeline
from src.pipeline.single_file_pipeline import SingleFilePipeline
from src.pipeline.fusion import ScoreFusion
from src.retrieval.candidates import select_candidates
from src.document_processing.file_reader import read_document
from src.document_processing.highlight_exporter import export_highlighted_html,export_highlighted_docx
from evaluation.fit_fusion import validation_partition


class FixedFusion:
    def __init__(self,p=.9):self.p=p
    def predict(self,values):
        return dict(probability=self.p,threshold_used=.5,decision='suspicious' if self.p>=.5 else 'checked_negative',probability_calibrated=False)


def cascade(p=.9,judge=None):
    return CaTEXPipeline(scorers=[lambda a,b:{'similarity':0.0}]*3+[lambda a,b:{'probability':.7}],
        fusion=FixedFusion(p),judge=judge or (lambda *a,**kw:{'explanation':'Evidence only.','label':'override','changed_phrases':[]}))


def pipeline(mode='full',p=.9,search=None,fetch=None,judge=None):
    return SingleFilePipeline(mode=mode,classifier=lambda s:[{'probability':.1,'decision':'skip'} for _ in s],
        search=search or (lambda *a,**kw:{'status':'ok','items':[{'url':'https://example.org/source'}]}),
        fetch=fetch or (lambda url:{'success':True,'text':'Original source sentence.'}),cascade=cascade(p,judge),
        semantic_scorer=lambda q,c:[.8]*len(c))


class PipelineTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.root=Path('tmp')/('pipeline-test-'+uuid4().hex);cls.root.mkdir(parents=True)

    def document(self,text='Short. Other sentence.'):
        path=self.root/'input.txt';path.write_text(text,encoding='utf-8');return read_document(path)

    def test_all_five_tiers_and_no_label_override(self):
        for p in (.1,.9):
            result=cascade(p).predict('Suspect.','Source.')
            self.assertEqual([x['tang'] for x in result['trace']],[1,2,3,4,5])
            self.assertEqual(result['decision'],'suspicious' if p>=.5 else 'checked_negative')
            self.assertNotEqual(result['label'],'override')

    def test_judge_failure_retains_decision_and_marks_incomplete(self):
        def fail(*a,**kw):raise RuntimeError('private error')
        result=cascade(judge=fail).predict('Suspect.','Source.')
        self.assertEqual(result['decision'],'suspicious');self.assertFalse(result['complete'])
        self.assertEqual(result['trace'][-1]['status'],'error')
        self.assertNotIn('private error',str(result))

    def test_modes_and_negative_counts(self):
        doc=self.document()
        full=pipeline(p=.1).process_document(doc)
        self.assertEqual(full['checked_sentences'],2);self.assertEqual(full['suspicious_sentences'],0)
        economy=pipeline(mode='economy').process_document(doc)
        self.assertEqual(economy['status_counts'],{'skipped_budget':2})
        self.assertEqual(economy['checked_sentences'],0)

    def test_distinct_search_extraction_model_errors(self):
        doc=self.document('Hello.')
        for response,status in [({'status':'no_results','items':[]},'no_source'),({'status':'error','items':[],'error':'timeout'},'search_failed')]:
            self.assertEqual(pipeline(search=lambda *a,**kw:response).process_document(doc)['results'][0]['status'],status)
        failed=pipeline(fetch=lambda u:{'success':False,'error':'read failed'}).process_document(doc)
        self.assertEqual(failed['results'][0]['status'],'extraction_failed')
        obj=pipeline();obj.cascade.predict=lambda *a:(_ for _ in ()).throw(ValueError('bad'))
        self.assertEqual(obj.process_document(doc)['results'][0]['status'],'model_failed')

    def test_partial_source_failure_is_not_negative(self):
        result=pipeline(p=.1,search=lambda *a,**k:{'status':'ok','items':[{'url':'https://a.test'},{'url':'https://b.test'}]},
            fetch=lambda u:{'success':True,'text':'Original source.'} if 'a.test' in u else {'success':False,'error':'timeout'}).process_document(self.document('Hello.'))
        self.assertEqual(result['results'][0]['status'],'incomplete')
        self.assertEqual(result['suspicious_sentences'],0)

    def test_hybrid_retains_semantic_only_and_no_arbitrary_zero(self):
        self.assertTrue(select_candidates('feline','A cat sleeps.',lambda q,c:[.9]*len(c)))
        self.assertEqual(select_candidates('feline','A cat sleeps.',lambda q,c:[0]*len(c)),[])

    def test_html_preserves_text_and_only_highlights_suspicious(self):
        from bs4 import BeautifulSoup
        result=pipeline().process_document(self.document('  Short.\n\n<script>danger</script> Other.'))
        result['results'][1]['status']='no_source';result['suspicious_sentences']=1
        path=export_highlighted_html(result,self.root/'report.html')
        soup=BeautifulSoup(path.read_text(encoding='utf-8'),'html.parser')
        self.assertEqual(soup.select_one('.document').get_text().lstrip(),result['original_text'].lstrip())
        self.assertIn('<div class="document">  <mark',path.read_text(encoding='utf-8'))
        self.assertEqual(len(soup.find_all('mark')),1);self.assertIsNone(soup.find('script'))

    def test_word_preserves_table_bold_and_selective_highlight(self):
        from docx import Document
        from docx.enum.text import WD_COLOR_INDEX
        path=self.root/'input.docx';doc=Document();p=doc.add_paragraph()
        p.add_run('First sentence. ').bold=True;p.add_run('Second sentence.').italic=True
        doc.add_table(rows=1,cols=1).cell(0,0).text='Table sentence.';doc.save(path)
        result=pipeline().process_document(read_document(path))
        result['results'][1]['status']='checked_negative';result['suspicious_sentences']=2
        output=export_highlighted_docx(result,self.root/'report.docx');actual=Document(output)
        self.assertEqual(actual.paragraphs[0].text,'First sentence. Second sentence.')
        yellow=''.join(r.text for r in actual.paragraphs[0].runs if r.font.highlight_color==WD_COLOR_INDEX.YELLOW)
        self.assertEqual(yellow,'First sentence.')
        self.assertTrue(actual.paragraphs[0].runs[0].bold)
        self.assertTrue(any(r.italic for r in actual.paragraphs[0].runs))
        self.assertEqual(actual.tables[0].cell(0,0).text,'Table sentence.')
        self.assertTrue(any(r.font.highlight_color==WD_COLOR_INDEX.YELLOW for r in actual.tables[0].cell(0,0).paragraphs[0].runs))

    def test_offsets_repeated_short_sentences_and_bad_encoding(self):
        result=pipeline().process_document(self.document('Yes. Yes.'))
        self.assertEqual([r['start'] for r in result['results']],[0,5])
        bad=self.root/'bad.txt';bad.write_bytes(b'\xff')
        with self.assertRaises(UnicodeDecodeError):read_document(bad)

    def test_fusion_group_partition_and_roundtrip(self):
        frame=pd.DataFrame([{'group_id':str(i),'candidate_group_id':str(j),'label':int(i==j)} for i in range(30) for j in range(30)])
        a,b=validation_partition(frame)
        self.assertFalse((set(a.group_id)|set(a.candidate_group_id)) & (set(b.group_id)|set(b.candidate_group_id)))
        x=np.array([[0,0,0,-2],[1,1,1,2]]*10);y=np.array([0,1]*10)
        fusion=ScoreFusion.fit(x,y,x,y,{'test':1});path=self.root/'fusion.pkl';fusion.save(path)
        loaded=ScoreFusion.load(path,{'test':1});self.assertEqual(loaded.predict(x[1])['decision'],'suspicious')
        with self.assertRaises(ValueError):ScoreFusion.load(path,{'test':2})

    def test_document_evaluation_keeps_abstentions_visible(self):
        from evaluation.evaluate_full_pipeline import evaluate_report
        result=pipeline(mode='economy').process_document(self.document('Hello.'))
        metrics=evaluate_report(result,pd.DataFrame({'sentence_id':[0],'label':[1]}))
        self.assertEqual(metrics['decision_coverage'],0)
        self.assertEqual(metrics['recall_over_all_labeled_positives'],0)
        self.assertIsNone(metrics['decided_subset_metrics'])

    def test_search_errors_do_not_leak_credentials(self):
        import requests
        from src.retrieval.web_search import search_web
        with patch.dict('os.environ',{'GOOGLE_SEARCH_API_KEY':'private-key','GOOGLE_SEARCH_ENGINE_ID':'cx'}):
            with patch('src.retrieval.web_search.requests.get',side_effect=requests.Timeout('private-key')):
                response=search_web('Some query')
            self.assertEqual(response['status'],'error')
            self.assertNotIn('private-key',str(response))
            with patch('src.retrieval.web_search.requests.get') as get:
                get.return_value.json.return_value={'items':[]}
                self.assertEqual(search_web('Some query')['status'],'no_results')

    def test_fusion_scoring_accepts_prepared_pair_schema(self):
        from evaluation.fit_fusion import score_pairs
        from contextlib import ExitStack
        frame=pd.DataFrame({'sentence_1':['Source.'],'sentence_2':['Suspect.']})
        with ExitStack() as stack:
            stack.enter_context(patch('src.models.tang1_tfidf.filter.load_vectorizer',return_value=object()))
            stack.enter_context(patch('src.models.tang2_sbert_pretrained.encode.load_model',return_value=(object(),object())))
            for name in ['src.models.tang1_tfidf.filter.tfidf_filter_batch','src.models.tang2_sbert_pretrained.encode.sbert_filter_batch','src.models.tang3_triplet_encoder.inference.tang3_predict_batch']:
                stack.enter_context(patch(name,return_value=pd.DataFrame({'similarity':[.2]})))
            last=stack.enter_context(patch('src.models.tang4_cross_encoder.inference.tang4_predict_batch',return_value=pd.DataFrame({'probability':[.5]})))
            scores=score_pairs(frame)
            self.assertEqual(last.call_args.args[0],[('Source.','Suspect.')])
            np.testing.assert_allclose(scores,[[.2,.2,.2,0.]])

    def test_source_extraction_and_single_class_metrics(self):
        from src.retrieval.extract_page_text import extract_content
        from evaluation.common import binary_metrics
        self.assertEqual(extract_content(b'<nav>Menu</nav><main><p>Evidence.</p><script>bad</script></main>','text/html'),'Evidence.')
        with self.assertRaises(ValueError):extract_content(b'unknown','application/octet-stream')
        report=binary_metrics([1,1],[.9,.1],[1,0])
        self.assertEqual(report['recall'],.5);self.assertIsNone(report['roc_auc'])

    def test_pair_order_matches_training_but_judge_keeps_public_order(self):
        seen=[]
        obj=cascade()
        obj.scorers=[lambda a,b:seen.append(('score',a,b)) or {'similarity':.2}]*3+[
            lambda a,b:seen.append(('cross',a,b)) or {'probability':.7}]
        obj.judge=lambda a,b,**kw:seen.append(('judge',a,b)) or {'explanation':'Evidence.','changed_phrases':[]}
        obj.predict('Suspect.','Source.')
        self.assertEqual(seen[-2],('cross','Source.','Suspect.'))
        self.assertEqual(seen[-1],('judge','Suspect.','Source.'))


if __name__=='__main__':unittest.main()
